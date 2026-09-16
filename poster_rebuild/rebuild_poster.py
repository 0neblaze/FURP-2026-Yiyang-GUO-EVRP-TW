"""Rebuild FURP poster: template as base, replace content precisely with python-pptx."""
import copy, shutil
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TPL = r"D:\UserData\Downloads\Poster Template - 副本.pptx"
OUT = r"D:\UserData\Documents\GitHub\FURP-2026-Yiyang-GUO-EVRP-TW\poster_rebuild\FURP_Showcase_Poster_v2.pptx"
ASSETS = r"D:\UserData\Documents\GitHub\FURP-2026-Yiyang-GUO-EVRP-TW\poster_rebuild\assets"

NOTT = RGBColor(0x10, 0x26, 0x3B)
BLUE = RGBColor(0x2E, 0x6C, 0xA4)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

shutil.copy(TPL, OUT)
prs = Presentation(OUT)
slide = prs.slides[0]

# ---------- helpers ----------
def walk_all(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk_all(sh.shapes)

def find_all(slide, name):
    return [sh for sh in walk_all(slide.shapes) if sh.name == name]

def find_one(slide, name, order=0, key="top"):
    lst = find_all(slide, name)
    lst.sort(key=lambda s: (s.top or 0, s.left or 0) if key == "top" else (s.left or 0, s.top or 0))
    return lst[order] if len(lst) > order else None

def set_runs(shape, lines, size=None, color=None, bold=None):
    """Replace text of a shape. lines: list[str]. First paragraph's first run formatting
    is reused for all lines; other paragraphs/runs are removed."""
    tf = shape.text_frame
    paras = tf.paragraphs
    p0 = paras[0]
    # ensure a run exists in p0
    if not p0.runs:
        r = p0.add_run()
        r.text = ""
    proto_rPr = None
    r0 = p0.runs[0]
    proto_rPr = copy.deepcopy(r0._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'))
    # remove all paragraphs except first; clear first
    txBody = tf._txBody
    a_ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for p in list(txBody.findall(a_ns + 'p'))[1:]:
        txBody.remove(p)
    for r in list(p0._p.findall(a_ns + 'r'))[1:]:
        p0._p.remove(r)
    r0.text = lines[0]
    if proto_rPr is not None:
        old = r0._r.find(a_ns + 'rPr')
        if old is not None:
            r0._r.remove(old)
        r0._r.insert(0, copy.deepcopy(proto_rPr))
    if size: r0.font.size = Pt(size)
    if color is not None: r0.font.color.rgb = color
    if bold is not None: r0.font.bold = bold
    # add remaining lines as new paragraphs cloned from p0
    for line in lines[1:]:
        new_p = copy.deepcopy(p0._p)
        txBody.append(new_p)
        from pptx.text.text import _Paragraph
        np = _Paragraph(new_p, tf)
        np.runs[0].text = line
    return shape

def add_text(slide, x, y, w, h, lines_specs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, space_after=0):
    """lines_specs: list of (text, font_name, size_pt, color, bold, line_spacing)."""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (text, fname, size, color, bold, lsp) in enumerate(lines_specs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if lsp: p.line_spacing = lsp
        if space_after: p.space_after = Pt(space_after)
        if text == "":
            continue
        # split symbol prefix if "SYM|rest"
        segs = []
        if "|" in text and text.split("|", 1)[0] in ("✓", "▲", "•"):
            sym, rest = text.split("|", 1)
            segs = [(sym, BLUE if sym == "✓" else NOTT), (rest, color)]
        else:
            segs = [(text, color)]
        for seg_text, seg_color in segs:
            r = p.add_run()
            r.text = seg_text
            r.font.name = fname
            r.font.size = Pt(size)
            r.font.color.rgb = seg_color
            r.font.bold = bold
    return tb

def add_pic(slide, path, x, y, w, h):
    return slide.shapes.add_picture(path, Cm(x), Cm(y), Cm(w), Cm(h))

def fit(img_w, img_h, box_w, box_h):
    """return (w, h) fitted inside box preserving ratio."""
    r = min(box_w / img_w, box_h / img_h)
    return img_w * r, img_h * r

# ---------- 1. delete placeholder Fig rects ----------
for nm in ["矩形 2", "矩形 4", "矩形 7", "矩形 14", "矩形 16", "矩形 22"]:
    sh = find_one(slide, nm)
    if sh is not None:
        sh._element.getparent().remove(sh._element)
        print("deleted placeholder:", nm)
    else:
        print("WARN not found:", nm)

# ---------- 2a. normalize KeyQuant group: remove 2.44x horizontal scale hack ----------
a_ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
p_ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
spTree = slide.shapes._spTree
fixed_keyquant = False
for grp in spTree.findall(p_ns + 'grpSp'):
    gpr = grp.find(p_ns + 'grpSpPr')
    if gpr is None: continue
    xfrm = gpr.find(a_ns + 'xfrm')
    if xfrm is None: continue
    off = xfrm.find(a_ns + 'off'); ext = xfrm.find(a_ns + 'ext'); chext = xfrm.find(a_ns + 'chExt')
    if off is None or ext is None or chext is None: continue
    if int(off.get('y')) == 21536605 and abs(int(off.get('x'))) <= 3:
        ext.set('cx', chext.get('cx'))   # scale_x 2.44 -> 1.0, box back to 24.41cm narrow
        fixed_keyquant = True
        print("KeyQuant group normalized: ext.cx -> %s (scale 1.0)" % chext.get('cx'))
assert fixed_keyquant, "KeyQuant group not normalized"

# ---------- 2b. right-bottom box: top-level copies of rect + banner (no group scaling) ----------
rect32 = ff56 = None
for grp in spTree.findall(p_ns + 'grpSp'):
    for sp in grp.findall(p_ns + 'sp'):
        cnv = sp.find('.//' + p_ns + 'cNvPr')
        if cnv is None: continue
        if cnv.get('name') == '矩形: 圆角 32': rect32 = sp
        if cnv.get('name') == '任意多边形: 形状 56': ff56 = sp
assert rect32 is not None and ff56 is not None, "source shapes for right-bottom box not found"
new_rect = copy.deepcopy(rect32)
new_ff = copy.deepcopy(ff56)
for el, (x, y, w, h) in [(new_rect, (25.66, 59.82, 33.78, 16.92)), (new_ff, (25.66, 59.82, 33.78, 2.66))]:
    xf = el.find(p_ns + 'spPr/' + a_ns + 'xfrm')
    o = xf.find(a_ns + 'off'); e = xf.find(a_ns + 'ext')
    o.set('x', str(int(round(x * 360000)))); o.set('y', str(int(round(y * 360000))))
    e.set('cx', str(int(round(w * 360000)))); e.set('cy', str(int(round(h * 360000))))
touched = False
for t in new_ff.iter(a_ns + 't'):
    if "Key Quantitative Highlights" in (t.text or ""):
        t.text = "Why It Works & Honest Audit"
        touched = True
assert touched, "cloned banner title not set"
spTree.append(new_rect)
spTree.append(new_ff)
print("right-bottom box added as top-level shapes (25.66,59.82) 33.78x16.92")

# ---------- 3. header title block ----------
tb4 = find_one(slide, "Text Box 4", key="left")  # (20.24,0)
tf = tb4.text_frame
a_ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
txBody = tf._txBody
paras = txBody.findall(a_ns + 'p')
# keep para0 (Title) and para1 (Name|Group) structure: rebuild cleanly
p_title = paras[0]
r_title = p_title.find('.//' + a_ns + 'r')
t_title = r_title.find(a_ns + 't')
t_title.text = "12.6% Fewer Vehicles, Zero Regressions:"
rPr = r_title.find(a_ns + 'rPr')
rPr.set('sz', "4800")
rPr.set('b', "1")
sf = rPr.find(a_ns + 'solidFill')
if sf is not None:
    rPr.remove(sf)
from lxml import etree
sf_new = etree.SubElement(rPr, a_ns + 'solidFill')
clr = etree.SubElement(sf_new, a_ns + 'srgbClr')
clr.set('val', "10263B")
# move solidFill to front of rPr (order matters in schema)
rPr.remove(sf_new)
rPr.insert(0, sf_new)
# add second title line by cloning paragraph
p2 = copy.deepcopy(p_title)
t2 = p2.find('.//' + a_ns + 't')
t2.text = "An Audit-First ALNS for Electric-Vehicle Routing"
txBody.insert(list(txBody).index(p_title) + 1, p2)
# subtitle line: find the paragraph containing "Name"
p_sub = None
for p in txBody.findall(a_ns + 'p'):
    full = "".join((t.text or "") for t in p.iter(a_ns + 't'))
    if full.startswith("Name"):
        p_sub = p
        break
if p_sub is not None:
    runs = p_sub.findall('.//' + a_ns + 'r')
    keep = runs[0]
    for r in runs[1:]:
        p_sub.remove(r)
    kt = keep.find(a_ns + 't')
    kt.text = "Yiyang Guo  |  Department of Mathematical Sciences"
    # keep original 32pt formatting from template
    print("header subtitle set")
print("header title set")

# ---------- 4. box banner texts stay; motivation icons & objective ----------
set_runs(find_one(slide, "文本框 1043"), ["E-VRPTW: battery-limited EVs, hard time windows"], size=20)
set_runs(find_one(slide, "文本框 1046"), ["Fleet size first in the lexicographic objective"], size=20)
set_runs(find_one(slide, "文本框 1047"), ["Feasibility hinges on exact charging decisions"], size=20)
obj = find_one(slide, "文本框 1053")
set_runs(obj, ["Objective — reproduce SSG (2014), then cut fleet size with zero regressions"], size=20)
obj.left = Cm(0.95)
obj.top = Cm(34.20)
obj.text_frame.margin_left = 0   # template default 0.25cm inset would misalign to 1.20
print("motivation texts set, objective aligned to x=0.95")

# ---------- 5. bottom bar KPI cards (text centered in each card) ----------
for nm in ["矩形: 圆角 1122", "矩形: 圆角 1127", "矩形: 圆角 1128"]:
    sh = find_one(slide, nm)
    assert sh is not None, nm + " not found"
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
set_runs(find_one(slide, "矩形: 圆角 1122"), ["−5.1% total distance"])
set_runs(find_one(slide, "矩形: 圆角 1127"), ["36/36 feasible formal runs"])
set_runs(find_one(slide, "矩形: 圆角 1128"), ["−22.0% exact-charging CPU"])
print("bottom bar cards set & centered")

# ---------- 6. footer ----------
ft_left = None
for sh in find_all(slide, "TextBox 1"):
    if sh.top and sh.top > Emu(int(80 * 360000)) and (sh.left or 0) < Emu(int(5 * 360000)):
        ft_left = sh
if ft_left is not None:
    # template box carries algn="ctr" + default insets that defeat precise alignment;
    # replace it with a clean textbox (margin 0, left-aligned, same font/size/color)
    ft_left._element.getparent().remove(ft_left._element)
add_text(slide, 0.95, 82.05, 26.0, 1.5,
         [("Department of Mathematical Sciences", "Roboto Condensed ExtraBold", 28, WHITE, False, 1.0)])
print("footer left rebuilt at x=0.95")
ft_right = None
for sh in find_all(slide, "TextBox 1"):
    if sh.top and sh.top > Emu(int(80 * 360000)) and (sh.left or 0) > Emu(int(20 * 360000)):
        ft_right = sh
if ft_right is not None:
    set_runs(ft_right, ["Yiyang Guo  (Student ID: 20719106)", "Supervisor: Tianxing Cui"])
    ft_right.left = Cm(27.00)
    ft_right.width = Cm(31.76)   # right edge 58.76 -> text right edge ≈ 58.5 (matches QR)
    print("footer right set")

# ---------- 7. figures (column margins unified: left col 0.95, right col 26.70) ----------
# Fig1: motivation box, x aligned 0.95, fills 16.9 width ; img 1584x1000
w, h = fit(1584, 1000, 16.9, 10.8)
add_pic(slide, ASSETS + r"\fig1_problem.png", 0.95 + (16.9 - w) / 2, 23.00, w, h)
print("fig1", w, h)
# benchmark / model / toolchain facts under the objective
# (catalog + AGENTS.md + review-report provenance; trimmed to 2 lines each to fit box)
GREY = RGBColor(0x44, 0x54, 0x6A)
add_text(slide, 0.95, 36.30, 16.9, 4.4, [
    ("Benchmark — SSG (2014) · 12 instances · 5–100 customers · C / R / RC",
     "Calibri", 15, GREY, False, 1.2),
    ("Model — battery-limited EVs · en-route charging · hard time windows",
     "Calibri", 15, GREY, False, 1.2),
], space_after=4)
add_text(slide, 0.95, 38.60, 16.9, 1.3, [
    ("Toolchain — Python · numba · CPLEX · Gurobi · HiGHS · OR-Tools (lockfile-pinned)",
     "Calibri", 12, GREY, False, 1.2)])
print("benchmark facts added")
# Fig2: method box left area (20.55,13.6) max 25.0 x 23.54 ; img 1166x1098
w, h = fit(1166, 1098, 25.0, 23.54)
add_pic(slide, ASSETS + r"\fig2_alns.png", 20.55, 13.60, w, h)
print("fig2", w, h)
# Fig3: exp-settings, x aligned 0.95 ; img 1506x1098
w, h = fit(1506, 1098, 16.9, 12.3)
add_pic(slide, ASSETS + r"\fig3_protocol.png", 0.95, 45.00, w, h)
print("fig3", w, h)
# Fig4: key results, left-aligned to 26.70 (right-column margin) ; img 2194x944
w, h = fit(2194, 944, 31.8, 12.4)
add_pic(slide, ASSETS + r"\fig4_fleet.png", 26.70, 44.75, w, h)
print("fig4", w, h)
# Fig5: key-quant box, x aligned 0.95, width 23.46 ; img 2116x862
w, h = fit(2116, 862, 23.46, 9.56)
add_pic(slide, ASSETS + r"\fig5_runtime.png", 0.95, 63.35, w, h)
print("fig5", w, h)
# Fig6 (NEW): per-instance vehicle savings chart, Method right column bottom
w, h = fit(2880, 1656, 12.20, 7.00)
add_pic(slide, ASSETS + r"\fig6_savings.png", 46.30, 32.10, w, h)
print("fig6 savings chart", w, h)
# QR: bottom-right of audit box, right edge 58.5 aligning with footer right margin
add_pic(slide, ASSETS + r"\qr_code.png", 54.30, 70.55, 4.2, 4.2)
print("QR placed")

# ---------- 8. method right column: two sections, 17pt, no duplication with bottom-right box ----------
col_x, col_w = 46.3, 12.2
add_text(slide, col_x, 14.0, col_w, 1.0,
         [("KEY MECHANISMS", "Roboto Condensed ExtraBold", 20, NOTT, True, 1.0)])
mech = [
    "•|Vehicle-count-aware repair: eliminate routes first; fleet can fall, never rise",
    "•|Optimistic prefilters (capacity · time · energy) screen before exact evaluation",
    "•|Constraint-guided destroy: station pressure, TW conflicts, energy detours",
    "•|Exact charging subproblems re-solve only the routes that changed",
]
add_text(slide, col_x, 15.5, col_w, 8.6,
         [(t, "Calibri", 17, BLACK, False, 1.15) for t in mech], space_after=7)
add_text(slide, col_x, 25.2, col_w, 1.0,
         [("AUDIT PROTOCOL", "Roboto Condensed ExtraBold", 20, NOTT, True, 1.0)])
audit = [
    "•|Lexicographic objective — vehicles first; acceptance rejects any fleet-size increase",
    "•|Independent review CLI replays solutions, event logs & hashes before publication",
]
add_text(slide, col_x, 26.7, col_w, 6.2,
         [(t, "Calibri", 17, BLACK, False, 1.15) for t in audit], space_after=7)
# chart section header + note (chart image itself placed in §7 as fig6)
add_text(slide, col_x, 30.95, col_w, 1.0,
         [("VEHICLE SAVINGS", "Roboto Condensed ExtraBold", 20, NOTT, True, 1.0)])
add_text(slide, col_x, 39.30, col_w, 1.2,
         [("Best of 3 seeds per instance — total fleet 87 → 76 vehicles (−12.6%)",
           "Calibri", 12, GREY, False, 1.15)])
print("method two-section column + savings chart section added")

# ---------- 9a. experimental settings: parameter card (right of Fig3) ----------
from pptx.enum.shapes import MSO_SHAPE
CARD_X, CARD_W = 18.15, 5.31          # right edge 23.46 = box inner right margin
card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(CARD_X), Cm(44.95), Cm(CARD_W), Cm(12.35))
card.adjustments[0] = 0.055
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(0xEF, 0xF5, 0xFD)
card.line.fill.background()
card.shadow.inherit = False
params = [("Instances", "12"), ("Seeds", "3"), ("Time limit", "30 s"), ("Iterations", "1,000"), ("Thread", "1")]
py = 45.45
for label, val in params:
    add_text(slide, CARD_X, py, CARD_W, 0.55, [(label, "Calibri", 13, NOTT, False, 1.0)],
             align=PP_ALIGN.CENTER)
    add_text(slide, CARD_X, py + 0.52, CARD_W, 1.05, [(val, "Roboto Condensed ExtraBold", 22, NOTT, True, 1.0)],
             align=PP_ALIGN.CENTER)
    py += 2.32
print("parameter card added")

# ---------- 9b. captions / notes lines (left col x=0.95, right col x=26.70) ----------
add_text(slide, 0.95, 57.40, 22.5, 0.9,
         [("Stage 0 frozen & hash-verified — every number replays from raw evidence",
           "Calibri", 15, BLACK, False, 1.0)])
add_text(slide, 26.70, 57.40, 31.8, 0.9,
         [("Vehicle-first search cuts fleet on 5/12 instances — 87 → 76 vehicles, −5.1% distance, 36/36 feasible",
           "Calibri", 15, BLACK, False, 1.0)])
# per-family fleet detail (computed live from stage00 vs attempt16 per-run CSVs)
add_text(slide, 26.70, 58.10, 31.8, 0.75,
         [("Per-family best — C 22→19 · R 31→28 · RC 34→29 · mean fleet 90.7 → 77.0",
           "Calibri", 13.5, NOTT, True, 1.0)])
add_text(slide, 0.95, 73.15, 23.46, 0.9,
         [("−22.0% exact-charging CPU (median, C instances) — CPU-batch default after audit",
           "Calibri", 15, BLACK, False, 1.0)])
add_text(slide, 0.95, 74.15, 23.46, 1.6,
         [("Batch pilot covers 4 instances (speedups 1.00–1.32×) — full-scope acceleration remains a Stage 3 target",
           "Calibri", 12, GREY, False, 1.2)])
add_text(slide, 20.55, 37.45, 25.0, 1.0,
         [("Constraint-guided ALNS engine — vehicle-first search with exact charging subproblems",
           "Calibri", 16, NOTT, True, 1.0)], align=PP_ALIGN.CENTER)
# cost structure + shipped acceleration lanes (review report: 1,687–2,161 exact calls / 30 s run)
add_text(slide, 20.55, 38.75, 25.0, 1.9, [
    ("Exact charging dominates cost: 1,687–2,161 calls per 30 s run",
     "Calibri", 13.5, BLACK, False, 1.25),
    ("Acceleration shipped: incremental cache · exact-deadline · parallel control",
     "Calibri", 13.5, BLACK, False, 1.25),
])
print("captions added")

# ---------- 10. right-bottom box: two-column claims + EVIDENCE AT SCALE cards ----------
from pptx.enum.shapes import MSO_CONNECTOR
RX = 26.70                    # box inner left margin (right-column reference)
# left column: WHY IT WORKS
LX, LW = 26.70, 15.5
add_text(slide, LX, 62.85, LW, 0.8,
         [("✓|WHY IT WORKS", "Roboto Condensed ExtraBold", 17, NOTT, True, 1.0)])
why = [
    "✓|Vehicle-count-aware repair — fleet can fall, never rise",
    "✓|Constraint-guided destroy targets station pressure · TW conflicts · energy detours",
    "✓|Dynamic removal size escalates on stagnation (17,892 logged)",
    "✓|Stage-gated: two independent complete reruns must agree",
]
add_text(slide, LX, 63.80, LW, 6.0,
         [(t, "Calibri", 15, BLACK, False, 1.18) for t in why], space_after=5)
# right column: HONEST BOUNDARIES
HX, HW = 43.00, 15.9
add_text(slide, HX, 62.85, HW, 0.8,
         [("▲|HONEST BOUNDARIES", "Roboto Condensed ExtraBold", 17, NOTT, True, 1.0)])
honest = [
    "▲|92/124 published BKS entries are model-incompatible — no gap claims against them",
    "▲|C5 control: −0.12% median (−8.7%…+4.5%) — speedup is structure-dependent, not universal",
    "▲|9/36 runs overran the 30 s budget by ≤ 0.02 s — all kept, with events",
    "▲|No unattested baselines — every number replays from raw solution JSON + event logs",
]
add_text(slide, HX, 63.80, HW, 6.0,
         [(t, "Calibri", 15, BLACK, False, 1.18) for t in honest], space_after=5)
# EVIDENCE AT SCALE: four stat cards (all numbers verified from repo / review report)
add_text(slide, RX, 70.35, 25.0, 0.7,
         [("EVIDENCE AT SCALE", "Roboto Condensed ExtraBold", 15, NOTT, True, 1.0)])
stats = [
    ("12,121", "evidence files on disk"),
    ("200,731", "operator event rows replayed"),
    ("20", "real failure cases kept"),
    ("0", "synthetic cases in audit"),
]
CARD_W2, GAP2 = 6.15, 0.30
for i, (num, lab) in enumerate(stats):
    cx = 26.70 + i * (CARD_W2 + GAP2)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(cx), Cm(71.10), Cm(CARD_W2), Cm(2.30))
    card.adjustments[0] = 0.09
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEF, 0xF5, 0xFD)
    card.line.fill.background()
    card.shadow.inherit = False
    add_text(slide, cx, 71.28, CARD_W2, 0.95, [(num, "Roboto Condensed ExtraBold", 22, NOTT, True, 1.0)],
             align=PP_ALIGN.CENTER)
    add_text(slide, cx + 0.25, 72.30, CARD_W2 - 0.5, 1.0, [(lab, "Calibri", 11, GREY, False, 1.05)],
             align=PP_ALIGN.CENTER)
# hairline separator + references (QR sits at right, placed in §7)
sep = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(RX), Cm(74.05), Cm(52.20), Cm(74.05))
sep.line.color.rgb = RGBColor(0xC9, 0xD2, 0xDD)
sep.line.width = Pt(0.75)
sep.shadow.inherit = False
refs = [
    "[1] Schneider, Stenger & Goeke (2014) Transp. Sci. 48(4):500–520    [2] Pisinger & Ropke (2006) Transp. Res. B",
    "[3] Keskin & Çatay (2016) Transp. Res. E    [4] Hottung, Kwon & Tierney (2022) EJOR",
]
add_text(slide, RX, 74.35, 26.0, 2.0, [(t, "Calibri", 12, GREY, False, 1.15) for t in refs])
add_text(slide, 49.8, 75.05, 8.7, 0.8,
         [("Scan for code & evidence", "Calibri", 12, BLACK, True, 1.0)], align=PP_ALIGN.RIGHT)
print("right-bottom box rebuilt: two columns + evidence-at-scale cards")

prs.save(OUT)
print("SAVED:", OUT)
